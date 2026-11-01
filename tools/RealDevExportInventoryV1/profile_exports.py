# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Sungwoo Kang
"""Profile locked dev exports without writing paths, cell values, or case rows."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import sys
import warnings
from collections import Counter
from pathlib import Path

import openpyxl

MAX_BYTES = 64 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
CASE_ID = re.compile(r"^[0-9a-f]{64}$")


def numeric(value: object) -> bool:
    if value is None or isinstance(value, bool):
        return False
    try:
        return math.isfinite(float(value))
    except (ValueError, TypeError, OverflowError):
        return False


def validate_file(path: Path, root: Path) -> Path:
    absolute = path.absolute()
    try:
        parts = absolute.relative_to(root).parts
    except ValueError as error:
        raise ValueError("source is outside the authorized corpus") from error
    cursor = root
    for part in parts:
        cursor = cursor / part
        if os.stat(cursor, follow_symlinks=False).st_file_attributes & 0x400:
            raise ValueError("reparse point in source path")
    if not absolute.is_file():
        raise ValueError("source file is missing")
    return absolute


def profile_rows(rows, counts: Counter) -> None:
    for index, row in enumerate(rows):
        if index >= MAX_ROWS or len(row) > MAX_COLUMNS:
            raise ValueError("worksheet exceeds bounded dimensions")
        cells = [cell for cell in row if cell is not None and str(cell).strip()]
        if not cells:
            continue
        counts["nonempty_rows"] += 1
        counts["formula_cells"] += sum(isinstance(cell, str) and cell.startswith("=") for cell in cells)
        if sum(numeric(cell) for cell in cells) >= 2:
            counts["rows_with_two_or_more_numeric_cells"] += 1


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--manifest-sha256", required=True)
    args = parser.parse_args()
    project_root = Path.cwd().resolve(strict=True)
    manifest = args.manifest.resolve(strict=True)
    manifest.relative_to((project_root / "artifacts" / "real-dev-export-inventory-v1").resolve(strict=True))
    raw_manifest = manifest.read_bytes()
    if hashlib.sha256(raw_manifest).hexdigest() != args.manifest_sha256.lower():
        raise ValueError("private manifest checksum mismatch")
    matches = json.loads(raw_manifest)
    if not isinstance(matches, list) or len(matches) != 13:
        raise ValueError("expected thirteen unique exact-name dev matches")
    corpus = (project_root / "data" / "manual data").resolve(strict=True)
    seen = set()
    counts: Counter = Counter()
    fingerprints = []
    failures: Counter = Counter()
    for record in matches:
        if set(record) != {"case_id", "project_path", "export_path"} or not CASE_ID.fullmatch(record["case_id"]):
            raise ValueError("private manifest record is invalid")
        if record["case_id"] in seen:
            raise ValueError("duplicate development case")
        seen.add(record["case_id"])
        project = validate_file(Path(record["project_path"]), corpus)
        export = validate_file(Path(record["export_path"]), corpus)
        if (project.suffix.lower() != ".dig" or export.suffix.lower() not in {".csv", ".xlsx"} or
                project.parent != export.parent or project.stem.casefold() != export.stem.casefold()):
            raise ValueError("export does not match its selected dev project")
        try:
            data = export.read_bytes()
            if not 0 < len(data) <= MAX_BYTES:
                raise ValueError("export exceeds bounded size")
            fingerprints.append(hashlib.sha256(data).hexdigest())
            file_counts: Counter = Counter()
            if export.suffix.lower() == ".csv":
                text = data.decode("utf-8-sig")
                profile_rows(csv.reader(io.StringIO(text, newline="")), file_counts)
                file_counts["csv_files_parsed"] += 1
            else:
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                    try:
                        file_counts["xlsx_worksheets"] += len(workbook.worksheets)
                        for sheet in workbook.worksheets:
                            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                                raise ValueError("worksheet exceeds bounded dimensions")
                            profile_rows(sheet.iter_rows(values_only=True), file_counts)
                    finally:
                        workbook.close()
                file_counts["xlsx_files_parsed"] += 1
            if file_counts["rows_with_two_or_more_numeric_cells"] > 0:
                file_counts["files_with_numeric_pair_rows"] += 1
            counts.update(file_counts)
        except Exception as error:
            failures[type(error).__name__] += 1

    report = {
        "schema": "graphreader.real-dev-export-structure-profile.v1",
        "matched_development_projects": len(matches),
        "totals": dict(counts),
        "failure_counts": dict(failures),
        "combined_export_sha256": hashlib.sha256("\n".join(sorted(fingerprints)).encode()).hexdigest(),
        "private_manifest_sha256": args.manifest_sha256.lower(),
        "profiler_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "read_only": True,
        "project_payloads_read": False,
        "sealed_payloads_read": False,
        "case_level_values_or_paths_in_report": False,
        "limitations": ["Exact-name matching does not prove semantic alignment.",
                        "Numeric rows are not verified graph-point or OCR ground truth."],
    }
    destination = manifest.parent / "structure-profile.json"
    with destination.open("x", encoding="utf-8", newline="\n") as output:
        json.dump(report, output, indent=2)
        output.write("\n")
    print(json.dumps(report, separators=(",", ":")))
    return 1 if failures else 0


if __name__ == "__main__":
    sys.dont_write_bytecode = True
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"REAL_DEV_EXPORT_PROFILE_FAILED {type(error).__name__}", file=sys.stderr)
        raise SystemExit(1) from None
