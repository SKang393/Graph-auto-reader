# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Sungwoo Kang
"""Compare locked dev Engauge points with matching exports, in aggregate only."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import sys
import warnings
from collections import Counter
from pathlib import Path

import openpyxl

from profile_exports import CASE_ID, MAX_BYTES, MAX_COLUMNS, MAX_ROWS, numeric, validate_file


def matching_point(x: float, y: float, rows: list[tuple[float, list[float]]], tolerance: float) -> bool:
    return any(math.isclose(x, row_x, rel_tol=1e-6, abs_tol=tolerance) and
               any(math.isclose(y, row_y, rel_tol=1e-6, abs_tol=tolerance) for row_y in values)
               for row_x, values in rows)


def unique_point_matches(points: list[dict], rows: list[tuple[float, list[float]]], tolerance: float) -> int:
    adjacency = [[index for index, row in enumerate(rows)
                  if matching_point(float(point["x"]), float(point["y"]), [row], tolerance)]
                 for point in points]
    row_to_point: dict[int, int] = {}

    def assign(point_index: int, visited: set[int]) -> bool:
        for row_index in adjacency[point_index]:
            if row_index in visited:
                continue
            visited.add(row_index)
            if row_index not in row_to_point or assign(row_to_point[row_index], visited):
                row_to_point[row_index] = point_index
                return True
        return False

    return sum(assign(index, set()) for index in range(len(points)))


def unmatched_category(x: float, y: float, rows: list[tuple[float, list[float]]]) -> str:
    x_present = any(math.isclose(x, row_x, rel_tol=1e-6, abs_tol=0.00501) for row_x, _ in rows)
    y_present = any(math.isclose(y, value, rel_tol=1e-6, abs_tol=0.00501)
                    for _, values in rows for value in values)
    if x_present and y_present:
        return "unmatched_x_and_y_in_different_rows"
    if x_present:
        return "unmatched_x_only"
    if y_present:
        return "unmatched_y_only"
    return "unmatched_neither_value"


def collect_rows(rows, result: list[tuple[float, list[float]]]) -> None:
    for index, row in enumerate(rows):
        if index >= MAX_ROWS or len(row) > MAX_COLUMNS:
            raise ValueError("worksheet exceeds bounded dimensions")
        if len(row) >= 2 and numeric(row[0]):
            values = [float(cell) for cell in row[1:] if numeric(cell)]
            if values:
                result.append((float(row[0]), values))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--private-points", type=Path)
    parser.add_argument("--private-points-sha256")
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        example: list[tuple[float, list[float]]] = []
        collect_rows([["x", "y"], ["1", "2", ""], [3, None, 4]], example)
        if (len(example) != 2 or not matching_point(1, 2, example, 1e-6) or
                not matching_point(3, 4, example, 1e-6) or matching_point(1, 4, example, 1e-6) or
                not matching_point(1.004, 2, example, 0.00501) or
                unmatched_category(1, 4, example) != "unmatched_x_and_y_in_different_rows" or
                unmatched_category(1, 8, example) != "unmatched_x_only" or
                unique_point_matches([{"x": 1, "y": 2}, {"x": 1, "y": 2}], example, 1e-6) != 1):
            raise ValueError("synthetic point matching self-test failed")
        print("REAL_DEV_EXPORT_POINT_COMPARISON_SELF_TEST_PASS")
        return 0
    if args.private_points is None or args.private_points_sha256 is None:
        parser.error("private points and checksum are required")
    root = Path.cwd().resolve(strict=True)
    source = args.private_points.resolve(strict=True)
    source.relative_to((root / "artifacts" / "real-dev-export-inventory-v1").resolve(strict=True))
    source_bytes = source.read_bytes()
    if hashlib.sha256(source_bytes).hexdigest() != args.private_points_sha256.lower():
        raise ValueError("private point manifest checksum mismatch")
    cases = json.loads(source_bytes)
    if not isinstance(cases, list) or len(cases) != 13:
        raise ValueError("expected thirteen unique exact-name dev matches")
    corpus = (root / "data" / "manual data").resolve(strict=True)
    seen = set()
    totals: Counter = Counter()
    failures: Counter = Counter()
    fingerprints = []
    for case in cases:
        if (set(case) != {"case_id", "project_path", "export_path", "project_sha256", "points"} or
                not CASE_ID.fullmatch(case["case_id"]) or not CASE_ID.fullmatch(case["project_sha256"]) or
                case["case_id"] in seen):
            raise ValueError("invalid development case record")
        seen.add(case["case_id"])
        project = validate_file(Path(case["project_path"]), corpus)
        export = validate_file(Path(case["export_path"]), corpus)
        if (project.suffix.lower() != ".dig" or export.suffix.lower() not in {".csv", ".xlsx"} or
                project.parent != export.parent or project.stem.casefold() != export.stem.casefold()):
            raise ValueError("export does not match selected development project")
        try:
            content = export.read_bytes()
            if not 0 < len(content) <= MAX_BYTES:
                raise ValueError("export exceeds bounded size")
            rows: list[tuple[float, list[float]]] = []
            if export.suffix.lower() == ".csv":
                collect_rows(csv.reader(io.StringIO(content.decode("utf-8-sig"), newline="")), rows)
            else:
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
                    try:
                        for sheet in workbook.worksheets:
                            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                                raise ValueError("worksheet exceeds bounded dimensions")
                            collect_rows(sheet.iter_rows(values_only=True), rows)
                    finally:
                        workbook.close()
            points = case["points"]
            if not isinstance(points, list) or any(set(point) != {"x", "y"} or
                    not numeric(point["x"]) or not numeric(point["y"]) for point in points):
                raise ValueError("invalid digitizer point record")
            strict = sum(matching_point(float(p["x"]), float(p["y"]), rows, 1e-6) for p in points)
            rounded = sum(matching_point(float(p["x"]), float(p["y"]), rows, 0.00501) for p in points)
            strict_unique = unique_point_matches(points, rows, 1e-6)
            rounded_unique = unique_point_matches(points, rows, 0.00501)
            totals.update({"cases_compared": 1, "digitizer_points": len(points), "numeric_export_rows": len(rows),
                           "strict_point_matches": strict, "rounded_point_matches": rounded,
                           "strict_one_to_one_matches": strict_unique, "rounded_one_to_one_matches": rounded_unique,
                           "cases_all_points_rounded_match": int(bool(points) and rounded == len(points)),
                           "cases_all_points_unique_rounded_match": int(bool(points) and rounded_unique == len(points)),
                           "cases_zero_points_rounded_match": int(bool(points) and rounded == 0)})
            if points and rounded == 0:
                totals[f"cases_zero_rounded_{export.suffix.lower()[1:]}"] += 1
            for point in points:
                x, y = float(point["x"]), float(point["y"])
                if not matching_point(x, y, rows, 0.00501):
                    totals[unmatched_category(x, y, rows)] += 1
                    totals["unmatched_reversed_xy_match"] += int(matching_point(y, x, rows, 0.00501))
            fingerprints.append(case["project_sha256"] + hashlib.sha256(content).hexdigest())
        except Exception as error:
            failures[type(error).__name__] += 1
    report = {
        "schema": "graphreader.real-dev-export-point-comparison-v1",
        "matched_development_projects": len(cases),
        "totals": dict(totals),
        "failure_counts": dict(failures),
        "strict_tolerance": "absolute 1e-6 or relative 1e-6 on x and y",
        "rounded_tolerance": "absolute 0.00501 or relative 1e-6 on x and y",
        "matching_rule": "first spreadsheet column is x; any subsequent numeric cell in that row can be y",
        "paired_content_sha256": hashlib.sha256("\n".join(sorted(fingerprints)).encode()).hexdigest(),
        "private_points_sha256": args.private_points_sha256.lower(),
        "comparison_source_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "sealed_payloads_read": False,
        "case_level_values_or_paths_in_report": False,
        "limitations": ["Exact-name matching does not establish semantic alignment.",
                        "Exports and points share a manual source and are not independent ground truth.",
                        "First-column and rounding assumptions may not hold for every export."],
    }
    with (source.parent / "point-comparison-v4.json").open("x", encoding="utf-8", newline="\n") as output:
        json.dump(report, output, indent=2)
        output.write("\n")
    print(json.dumps(report, separators=(",", ":")))
    return 1 if failures else 0


if __name__ == "__main__":
    sys.dont_write_bytecode = True
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"REAL_DEV_EXPORT_COMPARISON_FAILED {type(error).__name__}", file=sys.stderr)
        raise SystemExit(1) from None
